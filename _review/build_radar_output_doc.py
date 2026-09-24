from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE


OUT = r"D:\downloads\X2100_project-main\X2100雷达板数据输出与发射信号说明.docx"

# compact_reference_guide tokens
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
HEADER_FILL = "E8EEF5"
LIGHT_FILL = "F4F6F9"
MUTED = "5B6573"
INK = "1F2937"
TABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_TOP_BOTTOM = 80
CELL_START_END = 120


def set_font(run, size=None, bold=None, color=None, italic=None, latin="Calibri", east_asia="Microsoft YaHei"):
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=CELL_TOP_BOTTOM, start=CELL_START_END, bottom=CELL_TOP_BOTTOM, end=CELL_START_END):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[idx]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def style_table(table, header=True, first_col_bold=False, font_size=9.5):
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.08
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    set_font(run, size=font_size, bold=(r_idx == 0 or (first_col_bold and c_idx == 0)), color=INK)
        if r_idx == 0 and header:
            for cell in row.cells:
                set_cell_shading(cell, HEADER_FILL)
            set_repeat_table_header(row)


def add_table(doc, headers, rows, widths_dxa, first_col_bold=False, font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, text in enumerate(headers):
        table.rows[0].cells[i].text = text
    for values in rows:
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cells[i].text = str(value)
    set_table_geometry(table, widths_dxa)
    style_table(table, header=True, first_col_bold=first_col_bold, font_size=font_size)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)
    return table


def add_code_block(doc, lines):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.12)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.08
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), LIGHT_FILL)
    p_pr.append(shd)
    for idx, line in enumerate(lines):
        run = p.add_run(line)
        set_font(run, size=9.2, color="263238", latin="Consolas", east_asia="Microsoft YaHei")
        if idx != len(lines) - 1:
            run.add_break()
    return p


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_font(r1, bold=True, color=INK)
        r2 = p.add_run(text[len(bold_prefix):])
        set_font(r2, color=INK)
    else:
        r = p.add_run(text)
        set_font(r, color=INK)
    return p


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_font(run, size=9, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    run2 = paragraph.add_run(" 页")
    set_font(run2, size=9, color=MUTED)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

# Styles: compact_reference_guide. CJK font is an explicit readability override.
styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25

for name, size, color, before, after in (
    ("Heading 1", 16, BLUE, 18, 10),
    ("Heading 2", 13, BLUE, 14, 7),
    ("Heading 3", 12, DARK_BLUE, 10, 5),
):
    st = styles[name]
    st.font.name = "Calibri"
    st._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    st._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

# Running header/footer
hp = section.header.paragraphs[0]
hp.text = "X2100 雷达技术说明"
hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
hp.paragraph_format.space_after = Pt(0)
for run in hp.runs:
    set_font(run, size=9, color=MUTED)
add_page_number(section.footer.paragraphs[0])

# memo_masthead opening
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(12)
p.paragraph_format.space_after = Pt(4)
r = p.add_run("技术说明")
set_font(r, size=10, bold=True, color=BLUE)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(0)
p.paragraph_format.space_after = Pt(5)
r = p.add_run("X2100 雷达板数据输出与发射信号说明")
set_font(r, size=24, bold=True, color="172B4D")

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(14)
r = p.add_run("基于当前工程固件配置整理 | 2026-08-14")
set_font(r, size=10.5, color=MUTED)

callout = doc.add_table(rows=1, cols=1)
callout.style = "Table Grid"
callout.rows[0].cells[0].text = (
    "核心结论：当前固件默认通过 Type-C USB CDC 输出处理后的检测点、航迹、自车速度和预警结果；"
    "射频工作方式为 76.5 GHz 附近的线性 FMCW，2TX/4RX，配置采用 BPM-MIMO。"
    "射频前端送给 X2100 的底层数据是 4 路接收通道的差拍信号原始 ADC。"
)
set_table_geometry(callout, [TABLE_WIDTH_DXA])
set_repeat_table_header(callout.rows[0])
set_cell_shading(callout.rows[0].cells[0], LIGHT_FILL)
for run in callout.rows[0].cells[0].paragraphs[0].runs:
    set_font(run, size=10.5, bold=True, color=DARK_BLUE)
callout.rows[0].cells[0].paragraphs[0].paragraph_format.line_spacing = 1.2
doc.add_paragraph().paragraph_format.space_after = Pt(0)

doc.add_heading("1. 板级对外输出", level=1)
add_body(doc, "当前编译开关决定了默认输出路径和数据类型。正常工作时，Type-C 接口输出二进制 TLV 数据包；蓝色 USB-TTL 对应 UART3，主要承担命令、应答和 OTA 传输。")

add_table(
    doc,
    ["接口", "当前配置", "输出内容", "说明"],
    [
        ["Type-C USB CDC", "启用", "处理后的 TLV 二进制帧", "默认的数据采集接口"],
        ["UART3 / USB-TTL", "115200, 8N1", "命令、应答、OTA", "当前配置不从物理 UART3 发送默认数据帧"],
        ["CAN", "关闭", "无默认输出", "USE_CAN_TRANS=0"],
        ["TF 卡", "关闭", "不自动保存", "CHEETAH_SAVE=0"],
    ],
    [1700, 1600, 3100, 2960],
    first_col_bold=True,
)

doc.add_heading("1.1 默认 TLV 内容", level=2)
add_table(
    doc,
    ["TLV类型", "数据类别", "主要字段"],
    [
        ["21", "检测点/点云", "距离、方位角、径向速度、X/Y坐标、功率、SNR、运动状态等"],
        ["22", "目标航迹", "航迹ID、位置、速度、航向、有效状态、目标尺寸等"],
        ["23", "自车速度", "自车纵向速度估计，定点缩放100倍"],
        ["24", "预警结果", "BSD、LCA/LCW、DOW、RCW、FCW等状态"],
    ],
    [1300, 1900, 6160],
    first_col_bold=True,
)

doc.add_heading("1.2 数据包帧结构", level=2)
add_code_block(doc, [
    "28 字节帧头",
    "  Magic Word: 02 01 04 03 06 05 08 07",
    "  version / totalPacketLen / platform / frameNumber / numTLVs",
    "TLV 1: type + length + payload",
    "TLV 2: type + length + payload",
    "...",
])
add_body(doc, "当 Flash 参数 outRawDataFlg=0 时，输出上述处理结果；设为1并重启后，可切换为原始 ADC TLV。但当前 USB 发送实现可能产生 partial send，因此连续原始 ADC 更适合保存到 TF 卡。")

doc.add_heading("2. 射频发射信号配置", level=1)
add_body(doc, "当前工程采用线性调频连续波（FMCW）Chirp，天线配置为单颗 MMIC、2 路发射、4 路接收，并启用 BPM-MIMO。本文只记录该配置，不展开 BPM-MIMO 的工作原理。")

add_table(
    doc,
    ["参数", "当前值", "说明"],
    [
        ["波形", "线性 FMCW Chirp", "连续发射线性调频 Chirp"],
        ["MIMO配置", "BPM-MIMO", "2TX × 4RX，形成8个虚拟通道"],
        ["中心频率参数", "76.5 GHz", "工程 wave_params.central_freq"],
        ["调频斜率", "19.531 MHz/μs", "19.531 × 10¹² Hz/s"],
        ["Chirp周期", "26 μs", "单个 Chirp 的周期参数"],
        ["每帧Chirp数", "128", "单子帧配置"],
        ["ADC采样率", "26.665 MSps", "4路RX同步采样"],
        ["每Chirp采样点", "506", "当前为 cheetah_128_512_config"],
        ["正常帧周期", "50.328 ms", "约19.9帧/秒"],
        ["原始数据保存周期", "249.84 ms", "约4帧/秒，用于降低存储压力"],
        ["虚拟阵元间距", "0.5 λ", "8个虚拟阵元按工程顺序重排"],
    ],
    [2300, 2300, 4760],
    first_col_bold=True,
)

doc.add_heading("3. 射频前端输出给 X2100 的数据", level=1)
add_body(doc, "射频前端不会直接向 X2100 输出距离、速度、角度或目标点。前端首先将接收的毫米波回波与本振 Chirp 混频，得到低频差拍信号，经模拟滤波、增益处理和 ADC 采样后，再通过 CSI 数据通路送入 X2100。")
add_code_block(doc, [
    "毫米波回波 → 混频/去斜 → 差拍中频信号 → 滤波与增益 → ADC采样",
    "             → 4路RX原始ADC → CSI → X2100内存",
])

doc.add_heading("3.1 单帧原始 ADC 规模", level=2)
add_table(
    doc,
    ["项目", "数值", "计算/含义"],
    [
        ["接收通道", "4 RX", "RX0～RX3"],
        ["每Chirp采样点", "506", "每个RX各506点"],
        ["每帧Chirp数", "128", "当前单子帧"],
        ["样本存储宽度", "16 bit", "其中高12位为有效ADC值"],
        ["Chirp头", "32字节", "位于每个Chirp数据之前"],
        ["每Chirp字节数", "4,080字节", "32 + 506 × 4 × 2"],
        ["每帧字节数", "522,240字节", "4,080 × 128"],
    ],
    [2400, 2100, 4860],
    first_col_bold=True,
)

doc.add_heading("3.2 原始数据排列", level=2)
add_code_block(doc, [
    "Chirp 0: [32B Header] [Sample0: RX0 RX1 RX2 RX3] ... [Sample505: RX0 RX1 RX2 RX3]",
    "Chirp 1: [32B Header] [Sample0: RX0 RX1 RX2 RX3] ... [Sample505: RX0 RX1 RX2 RX3]",
    "...",
    "Chirp 127: [32B Header] [506 × 4 个 uint16 ADC样本]",
])
add_body(doc, "每个 ADC 样本使用 uint16 保存，固件按下式转换为以零为中心的幅值：")
add_code_block(doc, ["adc_value = (raw_uint16 >> 4) - 2048"])

doc.add_heading("4. X2100 后续处理及最终结果", level=1)
add_body(doc, "X2100 对4路原始 ADC 进行距离、速度、角度和目标级处理，最终生成适合上位机或车辆系统使用的结果。")
add_code_block(doc, [
    "4路原始ADC",
    "  → 去直流与加窗",
    "  → 距离FFT",
    "  → 多普勒FFT",
    "  → MIMO解码与角度FFT",
    "  → CFAR检测",
    "  → 检测点（距离/速度/角度/功率）",
    "  → 坐标转换与目标跟踪",
    "  → TLV点云、航迹、自车速度和预警结果",
])

doc.add_heading("5. 采数与 MATLAB 分析选择", level=1)
add_table(
    doc,
    ["分析目标", "应采集的数据", "推荐接口/介质", "MATLAB处理"],
    [
        ["绘制目标位置、速度和轨迹", "处理后TLV", "Type-C USB CDC", "解析帧头和TLV，绘制X/Y散点与轨迹"],
        ["验证车辆预警逻辑", "航迹与预警TLV", "Type-C USB CDC", "统计BSD/RCW等状态与目标关系"],
        ["自行做距离/速度FFT", "原始ADC", "TF卡", "解析4RX × 506 × 128数据立方体"],
        ["自行做角度估计", "原始ADC + 校准参数", "TF卡", "MIMO解码、通道校准和角度FFT"],
    ],
    [2200, 1900, 2000, 3260],
    first_col_bold=True,
    font_size=9.2,
)

doc.add_heading("6. 当前源码开关与核对方法", level=1)
add_code_block(doc, [
    "USE_USB_UART   = 1",
    "USE_USB_OUTPUT = 1",
    "UART_ADC_SEND  = 0",
    "SAVE_RAW_DATA  = 0",
    "CHEETAH_SAVE   = 0",
    "USE_CAN_TRANS  = 0",
])
add_body(doc, "上述结论基于当前工程源码。若板中烧录的是其他版本，实际输出可能不同。上电日志中应重点核对 outRawDataFlg、Total ADC size per frame、USB serial partial send、TF卡挂载与保存提示。")

doc.add_heading("参考源码位置", level=1)
refs = [
    r"firmware/x2100/freertos/vendor/motor_cycle_demo/src/set_params.c",
    r"firmware/x2100/freertos/vendor/motor_cycle_demo/inc/radar_types.h",
    r"firmware/x2100/freertos/vendor/motor_cycle_demo/src/get_adc_from_dat_file.c",
    r"firmware/x2100/freertos/vendor/mmw_msg_pkt/mmw_msg_pkt.h",
    r"firmware/x2100/freertos/vendor/mmw_msg_pkt/mmw_msg_pkt.c",
    r"firmware/x2100/freertos/vendor/uart_trans/uart_trans.c",
    r"firmware/x2100/freertos/vendor/vendor.c",
]
for ref in refs:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.375)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(ref)
    set_font(r, size=9.2, color=MUTED, latin="Consolas", east_asia="Microsoft YaHei")

# Metadata
doc.core_properties.title = "X2100 雷达板数据输出与发射信号说明"
doc.core_properties.subject = "当前固件的数据输出、射频参数和前端ADC格式"
doc.core_properties.author = "Codex"
doc.core_properties.keywords = "X2100, FMCW, ADC, TLV, radar"

doc.save(OUT)
print(OUT)
